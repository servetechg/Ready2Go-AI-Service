"""Document text extraction — swappable parser backend.

Two actors:
  DocumentParser  — protocol (interface) any backend must satisfy.
  BasicParser     — v1 default: pure-Python readers, no system deps.
                    PDF: pdfplumber (->pypdf fallback)
                    DOCX: python-docx
                    XLSX: openpyxl (sheet -> flat rows)
                    CSV: stdlib csv

Future backend (D5): LiteParseParser (adds OCR for scan-only PDFs).
Callers always receive the same Extraction dataclass regardless of backend.
"""

from __future__ import annotations

import csv
import io
import logging
from dataclasses import dataclass
from typing import Protocol, runtime_checkable

from app.config import get_settings

logger = logging.getLogger(__name__)

# A file with fewer chars than this is considered "empty".
_MIN_CHARS = 30


# ---------------------------------------------------------------------------
# Return types
# ---------------------------------------------------------------------------

@dataclass
class ExtractionQuality:
    """Deterministic quality read — drives the quality signal and hard overrides.

    is_scan_only: a PDF that produced pages but virtually no text (image/scan PDF).
    is_empty: any file that produced virtually no text at all.
    These flags flow directly into the scoring hard-overrides.
    """
    chars: int
    pages_or_rows: int
    is_scan_only: bool
    is_empty: bool


@dataclass
class Extraction:
    """Result of parsing one uploaded file."""
    text: str
    quality: ExtractionQuality


# ---------------------------------------------------------------------------
# Protocol — the contract every parser backend must satisfy
# ---------------------------------------------------------------------------

@runtime_checkable
class DocumentParser(Protocol):
    """Interface any parser backend must satisfy.

    Callers go through get_parser() so the backend can be swapped via config
    without any changes to callers (extract.py, api/integrity.py, etc.).
    """

    def extract(self, data: bytes, ext: str) -> Extraction:
        ...


# ---------------------------------------------------------------------------
# BasicParser — v1, pure Python, no system deps
# ---------------------------------------------------------------------------

class BasicParser:
    """Pure-Python text extraction for PDF/DOCX/XLSX/CSV.

    Design:
    - No length cap: reads the whole document (fixes the 8 000-char truncation
      of the old Next.js pipeline).
    - Scan-only detection: a PDF with pages but ~0 chars is an image PDF;
      flagged so the scorer applies the quality penalty rather than guessing.
    - Errors caught per-type: a broken PDF doesn't crash DOCX/CSV paths.
    """

    def extract(self, data: bytes, ext: str) -> Extraction:
        e = ext.lower().lstrip(".")
        try:
            if e == "pdf":
                return self._pdf(data)
            if e == "docx":
                return self._docx(data)
            if e in ("xlsx", "xls"):
                return self._xlsx(data)
            if e == "csv":
                return self._csv(data)
        except Exception:
            logger.exception("BasicParser.extract failed for ext=%s", ext)
        return _empty_extraction()

    # ---- per-type private methods ----------------------------------------

    def _pdf(self, data: bytes) -> Extraction:
        # pdfplumber first (better layout handling); pypdf as fallback.
        text, pages = _pdf_pdfplumber(data)
        if not text and pages > 0:
            # Got pages but no text — try pypdf as a second attempt.
            text2, _ = _pdf_pypdf(data)
            if text2:
                text = text2
        chars = len(text)
        is_scan = pages > 0 and chars < _MIN_CHARS
        return Extraction(
            text=text,
            quality=ExtractionQuality(
                chars=chars,
                pages_or_rows=pages,
                is_scan_only=is_scan,
                is_empty=chars < _MIN_CHARS,
            ),
        )

    def _docx(self, data: bytes) -> Extraction:
        import docx  # python-docx
        doc = docx.Document(io.BytesIO(data))
        lines = [p.text for p in doc.paragraphs if p.text.strip()]
        text = "\n".join(lines)
        chars = len(text)
        return Extraction(
            text=text,
            quality=ExtractionQuality(
                chars=chars,
                pages_or_rows=len(lines),
                is_scan_only=False,
                is_empty=chars < _MIN_CHARS,
            ),
        )

    def _xlsx(self, data: bytes) -> Extraction:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        rows_total = 0
        parts: list[str] = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            parts.append(f"--- Sheet: {sheet_name} ---")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    parts.append("\t".join(cells))
                    rows_total += 1
        wb.close()
        text = "\n".join(parts)
        chars = len(text)
        return Extraction(
            text=text,
            quality=ExtractionQuality(
                chars=chars,
                pages_or_rows=rows_total,
                is_scan_only=False,
                is_empty=chars < _MIN_CHARS,
            ),
        )

    def _csv(self, data: bytes) -> Extraction:
        # UTF-8 first; fall back to latin-1 for files with extended chars.
        try:
            text_data = data.decode("utf-8")
        except UnicodeDecodeError:
            text_data = data.decode("latin-1", errors="replace")
        reader = csv.reader(io.StringIO(text_data))
        rows = ["\t".join(r) for r in reader if any(c.strip() for c in r)]
        text = "\n".join(rows)
        chars = len(text)
        return Extraction(
            text=text,
            quality=ExtractionQuality(
                chars=chars,
                pages_or_rows=len(rows),
                is_scan_only=False,
                is_empty=chars < _MIN_CHARS,
            ),
        )


# ---------------------------------------------------------------------------
# PDF helpers (shared by BasicParser._pdf)
# ---------------------------------------------------------------------------

def _pdf_pdfplumber(data: bytes) -> tuple[str, int]:
    """Return (text, page_count). Swallows per-page errors."""
    try:
        import pdfplumber  # type: ignore[import-untyped]
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            pages = len(pdf.pages)
            parts = []
            for page in pdf.pages:
                try:
                    t = page.extract_text() or ""
                    if t.strip():
                        parts.append(t)
                except Exception:
                    pass
            return "\n".join(parts), pages
    except Exception:
        return "", 0


def _pdf_pypdf(data: bytes) -> tuple[str, int]:
    try:
        import pypdf  # type: ignore[import-untyped]
        reader = pypdf.PdfReader(io.BytesIO(data))
        pages = len(reader.pages)
        parts = []
        for page in reader.pages:
            try:
                t = page.extract_text() or ""
                if t.strip():
                    parts.append(t)
            except Exception:
                pass
        return "\n".join(parts), pages
    except Exception:
        return "", 0


def _empty_extraction() -> Extraction:
    return Extraction(
        text="",
        quality=ExtractionQuality(chars=0, pages_or_rows=0, is_scan_only=False, is_empty=True),
    )


# ---------------------------------------------------------------------------
# Factory — callers use this, never instantiate a parser directly
# ---------------------------------------------------------------------------

_PARSER_BACKENDS: dict[str, type] = {
    "basic": BasicParser,
    # "liteparse": LiteParseParser,  # future: enable for OCR support
}

_parser_instance: DocumentParser | None = None


def get_parser() -> DocumentParser:
    """Return the configured parser backend (cached singleton).

    Reads `parser_backend` from settings if present; defaults to BasicParser.
    The returned object always satisfies the DocumentParser protocol.
    """
    global _parser_instance
    if _parser_instance is None:
        settings = get_settings()
        backend_name = getattr(settings, "parser_backend", "basic").lower()
        cls = _PARSER_BACKENDS.get(backend_name, BasicParser)
        _parser_instance = cls()
    return _parser_instance
